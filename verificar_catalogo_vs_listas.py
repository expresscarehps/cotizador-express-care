#!/usr/bin/env python3
"""
Verifica que el catálogo de llantas del cotizador coincida con las listas de precios de los proveedores.
Se corre SIEMPRE que Carlos manda una lista nueva y ANTES de entregar el dev.

Uso:
  python3 verificar_catalogo_vs_listas.py cotizador_dev.html --serv "Servillantas 24sept26.xlsx" --vega "Vega 9oct26.xlsx"

Revisa (sale con código 1 si hay cualquier diferencia):
  Servillantas (lista = solo llantas con existencia, columnas CVE_ART/DESCRIPCION/MARCA/CHIH/PAR/MAYOREO):
    · llantas del cotizador que YA NO vienen en la lista
    · llantas de la lista que NO están en el cotizador
    · precio distinto
    · marca o descripción distinta (mismo CVE)
  VEGA (columnas CLAVE/ARTICULO/U. MEDIDA/PRECIO NETO/EXISTENCIA; el cotizador carga solo existencia > 0 y precio numérico):
    · llantas del cotizador que ya no vienen o que quedaron con existencia 0
    · llantas con existencia > 0 y precio que faltan en el cotizador
    · precio distinto
    · llantas con existencia pero sin precio numérico ("NO ESTA"): se avisan, no se pueden cargar
"""
import argparse, json, re, sys
import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument('html'); ap.add_argument('--serv'); ap.add_argument('--vega')
a = ap.parse_args()
h = open(a.html, encoding='utf8').read()

def arreglo(nombre):
    i = h.index('var ' + nombre + '=') + len('var ' + nombre + '='); j = h.index(';\n', i)
    return json.loads(h[i:j])

num = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
problemas = 0
def reporta(titulo, lista, grave=True):
    global problemas
    print(('  ❌ ' if (lista and grave) else '  ✅ ' if not lista else '  ⚠️ ') + titulo + ': ' + str(len(lista)))
    for x in lista[:15]: print('       ', x)
    if len(lista) > 15: print('        ... y', len(lista) - 15, 'más')
    if lista and grave: problemas += 1

if a.serv:
    print('\nSERVILLANTAS —', a.serv)
    cat = {x['cve'].strip(): x for x in arreglo('DATA')}
    lista = {}
    for r in list(openpyxl.load_workbook(a.serv, read_only=True).active.iter_rows(values_only=True))[3:]:
        if r[0] and num(r[5]): lista[str(r[0]).strip()] = r
    print('  llantas en la lista:', len(lista), '· en el cotizador:', len(cat))
    reporta('en el cotizador y YA NO vienen en la lista', [f"{k} {x['marca']} {x['desc']}" for k, x in cat.items() if k not in lista])
    reporta('en la lista y NO están en el cotizador', [f'{k} {r[2]} {r[1]}' for k, r in lista.items() if k not in cat])
    reporta('precio distinto', [f"{k} cotizador {cat[k]['precio']} lista {r[5]}" for k, r in lista.items() if k in cat and abs(cat[k]['precio'] - r[5]) > 0.005])
    reporta('marca o descripción distinta', [f"{k} cotizador [{cat[k]['marca']} {cat[k]['desc']}] lista [{(r[2] or '').strip()} {r[1].strip()}]" for k, r in lista.items() if k in cat and (cat[k]['marca'].upper() != (r[2] or '').strip().upper().replace('ZWARHTZ', 'ZWARTHZ') or cat[k]['desc'].strip() != r[1].strip())])

if a.vega:
    print('\nVEGA —', a.vega)
    cat = {x['cve'].strip(): x for x in arreglo('VEGA')}
    todas = {}
    for r in list(openpyxl.load_workbook(a.vega, read_only=True).active.iter_rows(values_only=True))[1:]:
        if r[0]: todas[str(r[0]).strip()] = r
    esperadas = {k: r for k, r in todas.items() if num(r[3]) and (r[4] or 0) > 0}
    print('  renglones en la lista:', len(todas), '· con existencia y precio:', len(esperadas), '· en el cotizador:', len(cat))
    reporta('en el cotizador y ya no vienen en la lista', [k for k in cat if k not in todas])
    reporta('en el cotizador pero con existencia 0 en la lista', [k for k in cat if k in todas and (todas[k][4] or 0) <= 0])
    reporta('con existencia y precio en la lista que FALTAN en el cotizador', [f'{k} {r[1]} ${r[3]} ({r[4]} pzas)' for k, r in esperadas.items() if k not in cat])
    reporta('precio distinto', [f"{k} cotizador {cat[k]['costo']} lista {todas[k][3]}" for k in cat if k in todas and num(todas[k][3]) and abs(cat[k]['costo'] - todas[k][3]) > 0.005])
    reporta('con existencia pero SIN precio numérico (no se pueden cargar; avisar a Carlos)', [f'{k} {r[1]} = {r[3]} ({r[4]} pzas)' for k, r in todas.items() if (r[4] or 0) > 0 and not num(r[3])], grave=False)

print('\nRESULTADO:', 'TODO COINCIDE' if problemas == 0 else f'{problemas} revisión(es) con diferencias')
sys.exit(1 if problemas else 0)
